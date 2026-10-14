from docx import Document

src = r"C:\Users\dlatj\Documents\카카오톡 받은 파일\인하대_위탁연구개발_계약서(초안)_rev1_261006.docx"
out = r"C:\Users\dlatj\Desktop\study\SAM\sam3-master\인하대_위탁연구개발_계약서(초안)_rev1_261006_수정본.docx"
old = "검지율 99% (픽셀 64×64 이상 기준)"
new = "10초 내 검지율 90% (1080p 기준, 픽셀 96×96 이상)"

doc = Document(src)
count = 0

def replace_in_paragraph(paragraph):
    global count
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            count += 1

for p in doc.paragraphs:
    replace_in_paragraph(p)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                replace_in_paragraph(p)

if count == 0:
    raise SystemExit("replacement text not found")
doc.save(out)
print(f"updated {count} occurrences: {out}")
